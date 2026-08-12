"""Verifica el libro localizando las celdas por su etiqueta, no por fila fija."""
import warnings
warnings.filterwarnings("ignore")
import formulas
from openpyxl import load_workbook

RUTA = "/home/user/DANISANTELMO/Alcopalet_Modelo_Financiero.xlsx"

# --- 1. Evaluar con `formulas` ---
xl = formulas.ExcelModel().loads(RUTA).finish()
sol = xl.calculate()

vals, errores = {}, []
for k, v in sol.items():
    if "]" not in k:
        continue
    clave = k.split("]", 1)[1].replace("'", "").upper()
    try:
        val = v.value[0, 0]
    except Exception:
        continue
    vals[clave] = val
    if isinstance(val, str) and val.startswith("#"):
        errores.append((clave, val))

print(f"Celdas evaluadas .......... {len(vals)}")
print(f"ERRORES DE FÓRMULA ........ {len(errores)}")
for c, v in errores:
    print("   ", c, "->", v)

# --- 2. Localizar filas por etiqueta ---
wb = load_workbook(RUTA)
idx = {}
for ws in wb.worksheets:
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if isinstance(c.value, str) and c.value.strip():
            idx.setdefault(ws.title.upper(), {})[c.value.strip()] = c.row


def v(hoja, etiqueta, col="B"):
    fila = idx.get(hoja.upper(), {}).get(etiqueta)
    if fila is None:
        return None, None
    return vals.get(f"{hoja.upper()}!{col}{fila}"), f"{hoja}!{col}{fila}"


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


CHEQUEOS = [
    ("PyG2025", "MARGEN BRUTO", "B", 1_892_823, 2, "€"),
    ("PyG2025", "EBITDA", "B", 452_674, 2, "€"),
    ("PyG2025", "RESULTADO NETO DEL EJERCICIO", "B", 279_364, 2, "€"),
    ("PyG2025", "Tipo impositivo efectivo", "B", 0.25, 0.001, ""),
    ("PyG2025", "CAJA OPERATIVA APROXIMADA 2025", "B", -60_555, 5, "€"),
    ("Escandallo", "COSTE DE PRODUCCIÓN UNITARIO", "B", 12.37, 0.01, "€"),
    ("Escandallo", "COSTE DE PRODUCCIÓN UNITARIO", "C", 7.02, 0.01, "€"),
    ("Escandallo", "COSTE DE PRODUCCIÓN UNITARIO", "D", 10.52, 0.01, "€"),
    ("Escandallo", "MARGEN DE CONTRIBUCIÓN UNITARIO", "B", 2.79, 0.01, "€"),
    ("Escandallo", "MARGEN DE CONTRIBUCIÓN UNITARIO", "C", 0.43, 0.01, "€"),
    ("Escandallo", "MARGEN DE CONTRIBUCIÓN UNITARIO", "D", -0.52, 0.01, "€"),
    ("Escandallo", "Operarios implícitos — palet nuevo", "B", 10.01, 0.05, ""),
    ("Escandallo", "Operarios implícitos — palet usado", "B", 6.99, 0.05, ""),
    ("Capacidad", "CAPACIDAD ACTUAL TOTAL", "D", 215_380, 1, "palets"),
    ("Capacidad", "GRADO DE UTILIZACIÓN", "B", 0.810, 0.002, ""),
    ("Capacidad", "Capacidad ociosa disponible hoy", "B", 40_876, 1, "palets"),
    ("Capacidad", "Robot — capacidad anual (palets)", "B", 87_120, 1, "palets"),
    ("Capacidad", "DÉFICIT/SUPERÁVIT del robot frente a Persán", "B", -2_880, 1, "palets"),
    ("Capacidad", "Recorrido disponible (palets/año)", "B", 99_670, 1, "palets"),
    ("Persan", "MARGEN DE CONTRIBUCIÓN", "B", -46_800, 2, "€"),
    ("Persan", "RESULTADO ANTES DE COSTE DE CIRCULANTE", "B", -154_800, 2, "€"),
    ("Persan", "IMPACTO ANUAL TOTAL ESTIMADO", "B", -188_876, 200, "€"),
    ("Persan", "3. + coste del circulante (equilibrio real)", "B", 12.17, 0.02, "€"),
    ("Robot", "AHORRO TEÓRICO BRUTO (€/año)", "B", 180_880, 200, "€"),
    ("Robot", "Plazo de recuperación de la inversión (meses)", "B", 2.65, 0.05, "meses"),
    ("Robot", "RESULTADO REAL DEL ROBOT AISLADO (€/año)", "B", -9_618, 20, "€"),
]

print("\n" + "=" * 78)
print("COMPROBACIÓN DE CIFRAS CLAVE (celda localizada por etiqueta)")
print("=" * 78)
fallos = 0
for hoja, etiq, col, esperado, tol, unidad in CHEQUEOS:
    val, ref = v(hoja, etiq, col)
    n = num(val)
    if n is None:
        print(f"  ??  {hoja:11s} {etiq[:44]:44s} NO LOCALIZADA")
        fallos += 1
        continue
    ok = abs(n - esperado) <= tol
    fallos += 0 if ok else 1
    print(f"  {'OK ' if ok else '!! '} {hoja:11s} {etiq[:44]:44s} {ref:>16s} = "
          f"{n:>14,.2f} {unidad}")

print(f"\nDesviaciones: {fallos} de {len(CHEQUEOS)}")

# --- 3. Escenarios ---
print("\n" + "=" * 78)
print("ESCENARIOS — impacto anual y resultado neto proyectado")
print("=" * 78)
cols = {"B": "E0 Statu quo", "C": "E1 Persán+robot", "D": "E2 Dividendo",
        "E": "E3 Robot solo", "F": "E4 Renegociado"}
for etiqueta in ("IMPACTO ANUAL SOBRE EL RESULTADO (€)",
                 "RESULTADO NETO PROYECTADO (€)",
                 "NECESIDAD DE CAJA DEL PRIMER AÑO (€)"):
    print(f"\n{etiqueta}")
    for col, nombre in cols.items():
        val, ref = v("Escenarios", etiqueta, col)
        n = num(val)
        print(f"   {nombre:20s} {n:>16,.0f} €" if n is not None else f"   {nombre:20s} —")

# --- 4. Sensibilidad: fila de madera estable ---
print("\n" + "=" * 78)
print("SENSIBILIDAD — impacto anual según precio negociado (madera estable)")
print("=" * 78)
ws = wb["Sensibilidad"]
fila_p = next(r for r in range(1, 30)
              if ws.cell(r, 2).value == 10.00 and isinstance(ws.cell(r, 3).value, float))
fila_0 = next(r for r in range(fila_p, fila_p + 12) if ws.cell(r, 1).value == 0.00)
for c in range(2, 10):
    p = ws.cell(fila_p, c).value
    if p is None:
        break
    n = num(vals.get(f"SENSIBILIDAD!{chr(64+c)}{fila_0}"))
    if n is not None:
        marca = "  <-- equilibrio" if -20000 < n < 20000 else ""
        print(f"   {p:5.2f} €/ud -> {n:>14,.0f} €/año{marca}")
