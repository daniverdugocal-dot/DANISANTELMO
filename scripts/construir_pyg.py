"""
Alcopalet (DTI-1399) — Análisis de la cuenta de resultados 2025.

Diseño: se lee de arriba abajo. Cada bloque es
    TÍTULO  ->  para qué sirve (a todo lo ancho)  ->  los números
Nada de notas en columnas lejanas.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
TIT = Font(name=F, size=16, bold=True, color="1F3864")
SUB = Font(name=F, size=9, italic=True, color="595959")
SEC = Font(name=F, size=11, bold=True, color="FFFFFF")
CAB = Font(name=F, size=10, bold=True, color="FFFFFF")
ETI = Font(name=F, size=10)
AZUL = Font(name=F, size=10, color="0000FF")
NEG = Font(name=F, size=10)
TOT = Font(name=F, size=10, bold=True)
BIG = Font(name=F, size=13, bold=True, color="C00000")
EXP = Font(name=F, size=9, italic=True, color="404040")
ETQ_SI = Font(name=F, size=9, bold=True, color="006100")
ETQ_NO = Font(name=F, size=9, bold=True, color="9C5700")

FS = PatternFill("solid", fgColor="1F3864")
FC = PatternFill("solid", fgColor="4472C4")
FT = PatternFill("solid", fgColor="D9E2F3")
FE = PatternFill("solid", fgColor="F2F2F2")
FA = PatternFill("solid", fgColor="FCE4E4")
FV = PatternFill("solid", fgColor="E2EFDA")
FR = PatternFill("solid", fgColor="FFF2CC")

fina = Side(style="thin", color="BFBFBF")
BOR = Border(left=fina, right=fina, top=fina, bottom=fina)

EUR = '#,##0 "€";(#,##0) "€";"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
MUL = '0.00"x"'

wb = Workbook()
ws = wb.active
ws.title = "PyG 2025"
ws.sheet_view.showGridLines = False
ANCHOS = {"A": 50, "B": 17, "C": 12, "D": 40}
for col, w in ANCHOS.items():
    ws.column_dimensions[col].width = w
ULT = "D"

ws["A1"] = "ALCOPALET — Cuenta de resultados 2025"
ws["A1"].font = TIT
ws["A2"] = ("Caso DTI-1399 · Lydes 2026 · Se lee de arriba abajo: cada bloque lleva el título, "
            "para qué sirve, y después los números.  ·  Azul = dato del Anexo 1 · Negro = fórmula")
ws["A2"].font = SUB
ws.merge_cells("A2:D2")

f = 4


def barra(texto, etiqueta=None):
    """Título de sección en banda azul. `etiqueta` marca si va al informe."""
    global f
    ws.cell(f, 1, texto).font = SEC
    for c in range(1, 5):
        ws.cell(f, c).fill = FS
    if etiqueta:
        cel = ws.cell(f, 4, etiqueta)
        cel.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        cel.alignment = Alignment(horizontal="right")
    f += 1


def explicar(texto):
    """Línea de explicación a todo el ancho, justo bajo el título."""
    global f
    c = ws.cell(f, 1, texto)
    c.font = EXP
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=4)
    for col in range(1, 5):
        ws.cell(f, col).fill = FE
    ancho_total = sum(ANCHOS.values())
    ws.row_dimensions[f].height = 12.5 * (len(texto) // int(ancho_total * 1.05) + 1) + 4
    f += 1


def cabecera(cols):
    global f
    for i, t in enumerate(cols):
        c = ws.cell(f, 1 + i, t)
        c.font = CAB
        c.fill = FC
        c.alignment = Alignment(horizontal="center" if i else "left")
        c.border = BOR
    f += 1


def linea(etiqueta, valor, fmt=EUR, fuente=NEG, pct=None, nota=None,
          relleno=None, destacar=False):
    global f
    ws.cell(f, 1, etiqueta).font = TOT if (relleno or destacar) else ETI
    if valor != "":
        c = ws.cell(f, 2, valor)
        c.font = BIG if destacar else (TOT if relleno else fuente)
        c.number_format = fmt
    if pct:
        p = ws.cell(f, 3, pct)
        p.font = TOT if relleno else NEG
        p.number_format = PCT
    if nota:
        n = ws.cell(f, 4, nota)
        n.font = EXP
        n.alignment = Alignment(wrap_text=True, vertical="center")
    if relleno:
        for c2 in range(1, 5):
            ws.cell(f, c2).fill = relleno
    f += 1
    return f - 1


def hueco(n=1):
    global f
    f += n


# ══════════════════════════════════════════════ RESUMEN (primera pantalla)
barra("LOS TRES NÚMEROS QUE SALEN DE ESTA CUENTA")
explicar("Si solo te llevas tres cifras al informe, son estas. Todo lo demás de la hoja "
         "existe para poder defenderlas.")
r_res = f
hueco(3)   # se rellenan al final, cuando ya existen las filas de origen

# ══════════════════════════════════════════════ LA CUENTA
barra("1. LA CUENTA DE RESULTADOS (Anexo 1)")
explicar("Tal cual la da el caso. Los subtotales están calculados con fórmula, no copiados, "
         "para comprobar que cuadran con lo publicado.")
cabecera(["Concepto", "Importe 2025", "% s/ventas", ""])

r_ven = linea("VENTAS NETAS", 4585267, EUR, AZUL, None, None, FT)
ws.cell(r_ven, 3, f"=B{r_ven}/$B${r_ven}").number_format = PCT
ws.cell(r_ven, 3).font = TOT


def pyg(et, val, nota=None):
    r = linea(et, val, EUR, AZUL, f"=B{f}/$B${r_ven}", nota)
    return r


r_vex = pyg("Variación de existencias", 376062, "Ingreso contable, pero NO es caja")
r_apr = pyg("Aprovisionamientos", -3068506, "COMPRAS, no consumo → ver bloque 2")
r_mb = linea("MARGEN BRUTO", f"=SUM(B{r_ven}:B{r_apr})", EUR, TOT,
             f"=B{f}/$B${r_ven}", None, FT)
r_pp = pyg("Gastos de personal — Producción", -481660)
r_tr = pyg("Gastos de transporte", -268218, "2 camiones, ya al 100% → ver bloque 7")
r_pe = pyg("Gastos de personal — Estructura", -311292)
r_oc = pyg("Otros costes de explotación", -378979, "60% es palet nuevo; varían con el volumen")
r_ebda = linea("EBITDA", f"=B{r_mb}+SUM(B{r_pp}:B{r_oc})", EUR, TOT,
               f"=B{f}/$B${r_ven}", None, FT)
r_am = pyg("Amortización", -36143)
r_ebit = linea("EBIT", f"=B{r_ebda}+B{r_am}", EUR, TOT, f"=B{f}/$B${r_ven}", None, FT)
r_gf = pyg("Gastos financieros", -44046, "Al 7,5% → ver bloque 6")
r_bai = linea("RESULTADO ANTES DE IMPUESTOS", f"=B{r_ebit}+B{r_gf}", EUR, TOT,
              f"=B{f}/$B${r_ven}", None, FT)
r_imp = pyg("Impuestos", -93121)
r_rn = linea("RESULTADO NETO", f"=B{r_bai}+B{r_imp}", EUR, TOT,
             f"=B{f}/$B${r_ven}", None, FT)
hueco()

# ══════════════════════════════════════════════ 2. FIJO / VARIABLE
barra("2. SEPARAR COSTES FIJOS DE VARIABLES", "→ VA AL INFORME")
explicar("EL cálculo del caso. Persán es un pedido incremental, y un pedido incremental no se "
         "juzga con costes completos sino con margen de contribución: lo que deja cada venta "
         "una vez pagado lo que esa venta consume. Sin esta separación no se puede opinar del "
         "contrato. Fíjate en que al final reconstruye el EBITDA exacto del Anexo 1.")
cabecera(["Concepto", "Importe", "% s/ventas", ""])

r_cons = linea("Consumo real de materia prima", f"=-B{r_apr}-B{r_vex}", EUR, NEG,
               f"=B{f}/B{r_ven}", "Compras menos lo que quedó en almacén")
linea("Transporte", f"=-B{r_tr}", EUR, NEG, f"=B{f}/B{r_ven}")
linea("Otros costes de explotación", f"=-B{r_oc}", EUR, NEG, f"=B{f}/B{r_ven}")
r_cv = linea("TOTAL COSTES VARIABLES", f"=SUM(B{f-3}:B{f-1})", EUR, TOT,
             f"=B{f}/B{r_ven}", "Crecen si se fabrica más", FT)
r_mc = linea("MARGEN DE CONTRIBUCIÓN", f"=B{r_ven}-B{r_cv}", EUR, TOT,
             f"=B{f}/B{r_ven}", "Cada euro vendido deja 27 céntimos", FV)
linea("Personal de producción", f"=-B{r_pp}", EUR, NEG, f"=B{f}/B{r_ven}",
      "Fijo: Esteban no quiere despidos")
linea("Personal de estructura", f"=-B{r_pe}", EUR, NEG, f"=B{f}/B{r_ven}")
r_cf = linea("TOTAL COSTES FIJOS", f"=B{f-2}+B{f-1}", EUR, TOT, f"=B{f}/B{r_ven}",
             "Se pagan igual se fabrique o no", FT)
linea("EBITDA reconstruido", f"=B{r_mc}-B{r_cf}", EUR, TOT, f"=B{f}/B{r_ven}")
linea("Diferencia con el EBITDA del Anexo 1", f"=B{f-1}-B{r_ebda}", EUR, TOT, None,
      "Cuadra al euro: la estructura es indiscutible", FV)
hueco()

# ══════════════════════════════════════════════ 3. CAJA
barra("3. ¿CUÁNTA CAJA GENERÓ REALMENTE LA EMPRESA?", "→ VA AL INFORME")
explicar("Hay dos peticiones de dinero sobre la mesa: la casa de los fundadores y los 643.000 € "
         "que exige Persán el primer año. La respuesta no está en el beneficio, está aquí. "
         "Alcopalet compró 376.062 € más de madera de la que consumió: eso es beneficio "
         "contable que salió del banco y sigue en el almacén.")
cabecera(["Concepto", "Importe", "", ""])

linea("EBITDA", f"=B{r_ebda}")
linea("(−) Aumento de existencias", f"=-B{r_vex}", EUR, NEG, None, "Madera comprada, aún sin vender")
linea("(−) Gastos financieros", f"=B{r_gf}")
linea("(−) Impuestos", f"=B{r_imp}")
r_caja = linea("CAJA OPERATIVA APROXIMADA", f"=SUM(B{f-4}:B{f-1})", EUR, TOT, None,
               "NEGATIVA, con 279.364 € de beneficio", FA, destacar=True)
linea("Resultado neto contable", f"=B{r_rn}")
r_brecha = linea("BRECHA entre beneficio y caja", f"=B{r_rn}-B{r_caja}", EUR, TOT, None,
                 "El beneficio está en el almacén, no en el banco", FA)
linea("Aviso a declarar en el informe", "", EUR, NEG, None,
      "No hay balance en el caso: se ignoran clientes y proveedores. Es aproximación.")
hueco()

# ══════════════════════════════════════════════ 4. PUNTO MUERTO
barra("4. PUNTO MUERTO Y APALANCAMIENTO", "→ VA AL INFORME")
explicar("Mide el colchón: cuánto puede caer la facturación antes de entrar en pérdidas. Sirve "
         "para dos cosas opuestas y ambas útiles: demuestra que la empresa está sana HOY, y "
         "demuestra que Persán empeora las dos variables a la vez (sube los costes fijos en "
         "108.000 € de campa y leasing, y aporta margen negativo).")
cabecera(["Concepto", "Importe", "", ""])

r_ratio = linea("Ratio de margen de contribución", f"=B{r_mc}/B{r_ven}", PCT, TOT)
linea("Punto muerto (EBITDA = 0)", f"=B{r_cf}/B{r_ratio}", EUR, NEG, None, "Ventas para no perder caja")
r_pm = linea("PUNTO MUERTO (resultado = 0)", f"=(B{r_cf}-B{r_am}-B{r_gf})/B{r_ratio}",
             EUR, TOT, None, "Ventas para no entrar en pérdidas", FT)
linea("Ventas reales 2025", f"=B{r_ven}")
r_seg = linea("MARGEN DE SEGURIDAD", f"=(B{r_ven}-B{r_pm})/B{r_ven}", PCT, TOT, None,
              "Puede caer un 30% antes de perder dinero", FV, destacar=True)
linea("Apalancamiento operativo", f"=B{r_mc}/B{r_ebda}", MUL, TOT, None,
      "1% menos de margen mueve el EBITDA un 2,75%")
hueco()

# ══════════════════════════════════════════════ 5. COMPROBACIÓN ANEXO 3
barra("5. COMPROBACIÓN: ¿CUADRA EL ANEXO 1 CON EL ANEXO 3?", "no va al informe")
explicar("Esto no se publica: es lo que te autoriza a usar el escandallo del Anexo 3 sin que "
         "nadie te lo discuta. Si alguien dice que los anexos no cuadran, esta es la respuesta.")
cabecera(["Concepto", "Importe", "", ""])

linea("Consumo real según el Anexo 1", f"=B{r_cons}")
linea("Materia prima según el Anexo 3", 2692732, EUR, AZUL, None,
      "8,27 € × 174.504 nuevos + 4,80 € × 260.330 usados")
linea("DESVIACIÓN", f"=B{f-2}-B{f-1}", EUR, TOT, None,
      "288 € sobre 2,7 millones: los anexos son coherentes", FV)
hueco()

# ══════════════════════════════════════════════ 6. DEUDA
barra("6. ENDEUDAMIENTO IMPLÍCITO", "al anexo / turno de preguntas")
explicar("El caso no da balance, pero los gastos financieros divididos por el tipo del banco "
         "reconstruyen la deuda aproximada. Es tu respuesta si te preguntan «¿y no lo puede "
         "financiar el banco?».")
cabecera(["Concepto", "Importe", "", ""])

r_tipo = linea("Tipo de interés del banco", 0.075, PCT, AZUL)
r_deuda = linea("Deuda implícita", f"=-B{r_gf}/B{r_tipo}", EUR, TOT)
linea("Deuda / EBITDA — HOY", f"=B{r_deuda}/B{r_ebda}", MUL, TOT, None,
      "Por debajo de 2x: el banco presta sin problema", FV)
linea("Cobertura de intereses", f"=B{r_ebda}/-B{r_gf}", MUL, NEG)
linea("Deuda estimada TRAS Persán", f"=B{r_deuda}+190000+454356", EUR, NEG, None,
      "+190.000 leasing +454.356 circulante a 180 días")
linea("EBITDA estimado TRAS Persán", f"=B{r_ebda}-154800", EUR, NEG, None,
      "−46.800 margen −60.000 campa −48.000 leasing")
r_dpost = linea("Deuda / EBITDA — TRAS PERSÁN", f"=B{f-2}/B{f-1}", MUL, TOT, None,
                "Por encima de 4x el banco pone condiciones", FA, destacar=True)
hueco()

# ══════════════════════════════════════════════ 7. TRANSPORTE
barra("7. COSTE REAL DEL TRANSPORTE", "inferencia propia")
explicar("El contrato exige DOS camiones de retén permanente. Conviene comprobar si el coste "
         "que el Anexo 3 asigna a Persán es coherente con lo que cuestan los camiones actuales. "
         "OJO: es una inferencia tuya, no un dato del caso, y hay que presentarla como tal.")
cabecera(["Concepto", "Importe", "", ""])

r_ct = linea("Coste total de transporte 2025", f"=-B{r_tr}")
r_ncam = linea("Número de camiones actuales", 2, NUM, AZUL)
linea("COSTE POR CAMIÓN Y AÑO", f"=B{r_ct}/B{r_ncam}", EUR, TOT, None, None, FT)
linea("Transporte asignado a Persán (Anexo 3)", "=0.53*90000", EUR, NEG, None,
      "0,53 €/palet × 90.000 unidades")
linea("DIFERENCIA si el 3er camión va a Persán", f"=B{f-2}-B{f-1}", EUR, TOT, None,
      "El escandallo es optimista: el contrato es peor", FA)
hueco()

# ══════════════════════════════════════════════ CONCLUSIÓN
barra("QUÉ DICE, EN CUATRO FRASES")
for txt in [
    "El 72,8% de la facturación se va en costes variables, y la madera sola es el 58,7%. "
    "Alcopalet es, en lo económico, una transformadora de madera con margen estrecho sobre "
    "una materia prima que no controla.",
    "Con esa estructura el precio de venta no es una variable comercial: es LA variable. "
    "Vender a 10 € un palet cuyo coste variable es 10,52 € no es un descuento agresivo.",
    "El apalancamiento de 2,75x amplifica cualquier error de precio, y Persán añade 108.000 € "
    "de costes fijos sin aportar contribución alguna.",
    "La empresa está sana hoy (30% de margen de seguridad) pero no generó caja. El debate del "
    "comité está mal planteado: no es «invertir o repartir», es que ahora mismo no hay caja "
    "para ninguna de las dos cosas.",
]:
    c = ws.cell(f, 1, "• " + txt)
    c.font = ETI
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=f, start_column=1, end_row=f, end_column=4)
    ws.row_dimensions[f].height = 12.5 * (len(txt) // 118 + 1) + 6
    f += 1

# ══════════════════════════════════════════════ Rellenar el resumen de arriba
resumen = [
    ("Margen de contribución del negocio", f"=B{r_mc}/B{r_ven}", PCT,
     "Lo que deja cada euro vendido. El número clave para juzgar Persán."),
    ("Caja operativa real 2025", f"=B{r_caja}", EUR,
     "Frente a 279.364 € de beneficio. La respuesta a los padres."),
    ("Margen de seguridad", f"=B{r_seg}", PCT,
     "Cuánto puede caer la facturación antes de perder dinero."),
]
for i, (et, formula, fmt, nota) in enumerate(resumen):
    r = r_res + i
    ws.cell(r, 1, et).font = TOT
    c = ws.cell(r, 2, formula)
    c.font = BIG
    c.number_format = fmt
    n = ws.cell(r, 4, nota)
    n.font = EXP
    n.alignment = Alignment(wrap_text=True, vertical="center")
    for col in range(1, 5):
        ws.cell(r, col).fill = FR
        ws.cell(r, col).border = BOR

ws.freeze_panes = "A4"
wb.save("/home/user/DANISANTELMO/Alcopalet_PyG_2025.xlsx")
print(f"Guardado. Filas usadas: {f-1}")
