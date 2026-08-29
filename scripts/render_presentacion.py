"""
Renderiza la presentación a imágenes sin LibreOffice, para revisión visual.

Dibuja la geometría real del .pptx (posiciones, tamaños, rellenos y texto) con
Pillow. La tipografía es aproximada —sustituye Cambria y Calibri por DejaVu—,
así que sirve para comprobar composición, no para juzgar el tipo de letra.
"""
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

RUTA = sys.argv[1] if len(sys.argv) > 1 else "/home/user/DANISANTELMO/Alcopalet_Defensa.pptx"
SALIDA = Path(sys.argv[2] if len(sys.argv) > 2 else "/tmp/slides")
SALIDA.mkdir(exist_ok=True)
PPP = 110                      # píxeles por pulgada
EMU = 914400.0

SERIF = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"
SERIF_B = "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"
SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
SANS_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_cache = {}


def fuente(nombre, pt, negrita):
    serif = (nombre or "").startswith("Cambria")
    ruta = (SERIF_B if negrita else SERIF) if serif else (SANS_B if negrita else SANS)
    px = max(int(pt * PPP / 72 * 0.92), 7)
    clave = (ruta, px)
    if clave not in _cache:
        _cache[clave] = ImageFont.truetype(ruta, px)
    return _cache[clave]


def px(v):
    return int(round(v / EMU * PPP))


def color_de(obj, defecto=None):
    try:
        if obj and obj.type is not None and obj.rgb is not None:
            return "#" + str(obj.rgb)
    except Exception:
        pass
    return defecto


prs = Presentation(RUTA)
W, H = px(prs.slide_width), px(prs.slide_height)

for n, slide in enumerate(prs.slides, 1):
    fondo = "#FFFFFF"
    try:
        if slide.background.fill.type is not None:
            fondo = color_de(slide.background.fill.fore_color, "#FFFFFF")
    except Exception:
        pass
    im = Image.new("RGB", (W, H), fondo)
    d = ImageDraw.Draw(im)

    for sh in slide.shapes:
        if sh.left is None:
            continue
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)

        # relleno de la forma
        relleno = None
        try:
            if sh.fill.type is not None and sh.fill.type != 5:
                relleno = color_de(sh.fill.fore_color)
        except Exception:
            pass
        if relleno:
            forma = str(sh.shape_type)
            if "OVAL" in forma or "ELLIPSE" in forma:
                d.ellipse([x, y, x + w, y + h], fill=relleno)
            else:
                d.rounded_rectangle([x, y, x + w, y + h], radius=6, fill=relleno)
        if "LINE" in str(sh.shape_type) and not relleno:
            d.line([x, y, x + w, y + h], fill="#EFD9B8", width=2)

        if not sh.has_text_frame:
            continue

        cur_y = y + 2
        for par in sh.text_frame.paragraphs:
            texto = "".join(r.text for r in par.runs)
            if not texto.strip():
                cur_y += 6
                continue
            r0 = par.runs[0]
            pt = r0.font.size.pt if r0.font.size else 14
            nom = r0.font.name or "Calibri"
            neg = bool(r0.font.bold)
            col = color_de(r0.font.color, "#000000")
            f = fuente(nom, pt, neg)
            alin = str(par.alignment or "")

            # ajuste de línea al ancho de la caja
            palabras, linea, lineas = texto.split(" "), "", []
            for p in palabras:
                prueba = (linea + " " + p).strip()
                if d.textlength(prueba, font=f) <= w - 6 or not linea:
                    linea = prueba
                else:
                    lineas.append(linea)
                    linea = p
            lineas.append(linea)

            for ln in lineas:
                anchura = d.textlength(ln, font=f)
                if "CENTER" in alin:
                    tx = x + (w - anchura) / 2
                elif "RIGHT" in alin:
                    tx = x + w - anchura - 3
                else:
                    tx = x + 3
                d.text((tx, cur_y), ln, font=f, fill=col)
                cur_y += int(pt * PPP / 72 * 1.22)

    im.save(SALIDA / f"slide-{n:02d}.png")

print(f"{len(prs.slides._sldIdLst)} diapositivas en {SALIDA}")
